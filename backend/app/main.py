import asyncio
import csv
import io
import logging
import os
import re
import secrets
import warnings
import httpx
from collections import defaultdict, deque
from contextlib import asynccontextmanager, suppress
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Annotated
from zoneinfo import ZoneInfo
from dotenv import load_dotenv

from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Request, Response, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from PIL import Image, ImageOps, UnidentifiedImageError
from sqlalchemy import delete, func, select, text as sql_text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .db import Base, make_engine, session_factory
from .models import AIAssessment, AIReviewJob, Area, AuthSession, Brigade, Employee, Equipment, FaultCode, Material, MaterialWriteoff, Notification, Order, OrderEvent, Photo, TimeNorm, utcnow
from .ai import ANSWER_SCHEMA, HINT_SCHEMA, INSIGHTS_SCHEMA, OpenAIProvider, assistant_facts, evidence_facts, process_one_review, redact_personnel, supported_evidence
from .ai_service_client import AIServiceProvider
from .schemas import AssistantQuestion, Completion, Login, OrderCreate, OrderHintRequest, OrderPatch, Transition
from .security import check_pin, hash_pin, token_hash
from .seed import seed_database
from .services import AIReviewStub, EXECUTION_FINISHED, TERMINAL, analytics, audit, aware, downtime_minutes, employee_dict, iso, monitor_deadlines, notify, order_dict, photo_dict, shift_start

log = logging.getLogger(__name__)
STATUS = {"issued", "accepted", "queued", "rejected", "in_progress", "paused", "completed", "ai_review", "rework", "closed", "cancelled"}
PRIORITY = {"emergency", "high", "normal", "planned"}
WS_AUTH_RECHECK_SECONDS = 30


class Realtime:
    def __init__(self):
        self.clients = set()

    async def publish(self, type_, order_id=None):
        message = {"type": type_}
        if order_id:
            message["order_id"] = order_id
        for ws in list(self.clients):
            try:
                await asyncio.wait_for(ws.send_json(message), timeout=2)
            except Exception:
                self.clients.discard(ws)


def create_app(database_url=None, seed=True, monitor=True, ai_provider=None, ai_worker=True):
    load_dotenv(Path(__file__).resolve().parents[2] / ".env", override=False)
    if ai_provider is None and os.getenv("AI_PROVIDER", "stub").lower() == "openai":
        ai_provider = OpenAIProvider()
    elif ai_provider is None and os.getenv("AI_PROVIDER", "stub").lower() == "service":
        ai_provider = AIServiceProvider()
    engine = make_engine(database_url)
    sessions = session_factory(engine)
    realtime = Realtime()
    attempts = defaultdict(deque)

    def run_monitor():
        with sessions() as db:
            return monitor_deadlines(db)

    async def deadline_loop():
        while True:
            try:
                added = await asyncio.to_thread(run_monitor)
                if added:
                    await realtime.publish("notifications.updated")
            except asyncio.CancelledError:
                raise
            except Exception:
                log.exception("Deadline monitor failed; retrying in five seconds")
            await asyncio.sleep(5)

    async def ai_loop():
        while True:
            processed = False
            try:
                processed = await asyncio.to_thread(process_one_review, sessions, ai_provider)
                if processed:
                    await realtime.publish("orders.updated")
                    await realtime.publish("notifications.updated")
            except asyncio.CancelledError:
                raise
            except Exception:
                log.exception("AI review worker failed")
            await asyncio.sleep(2 if processed else 5)

    @asynccontextmanager
    async def lifespan(app):
        Base.metadata.create_all(engine)
        if seed and os.getenv("SEED_DEMO", "true").lower() == "true":
            with sessions() as db:
                seed_database(db)
        task = asyncio.create_task(deadline_loop()) if monitor else None
        ai_task = asyncio.create_task(ai_loop()) if ai_provider and ai_worker else None
        yield
        for running_task in (task, ai_task):
            if running_task:
                running_task.cancel()
                with suppress(asyncio.CancelledError):
                    await running_task
        for ws in list(realtime.clients):
            with suppress(Exception):
                await ws.close(code=1001)
        engine.dispose()

    app = FastAPI(title="НарядAI API", version="1.0.0", lifespan=lifespan)
    app.state.engine = engine
    app.state.sessions = sessions
    app.state.realtime = realtime
    app.add_middleware(CORSMiddleware, allow_origins=os.getenv("CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173,http://localhost:8080,http://127.0.0.1:8080").split(","), allow_credentials=False, allow_methods=["GET", "POST", "PATCH", "OPTIONS"], allow_headers=["Authorization", "Content-Type"])

    def get_db():
        with sessions() as db:
            try:
                yield db
            except Exception:
                db.rollback()
                raise

    DB = Annotated[Session, Depends(get_db)]

    def lookup_user(db, token):
        session = db.scalar(select(AuthSession).where(AuthSession.token_hash == token_hash(token)))
        if not session or aware(session.expires_at) <= utcnow():
            raise HTTPException(401, "Сессия истекла. Войдите снова.", headers={"WWW-Authenticate": "Bearer"})
        user = db.get(Employee, session.employee_id)
        if not user:
            raise HTTPException(401, "Пользователь не найден")
        return user

    def current_user(request: Request, db: DB):
        auth = request.headers.get("Authorization", "")
        if not auth.startswith("Bearer "):
            raise HTTPException(401, "Необходима авторизация", headers={"WWW-Authenticate": "Bearer"})
        return lookup_user(db, auth[7:])

    User = Annotated[Employee, Depends(current_user)]

    def require_role(user, *roles):
        if user.role not in roles:
            raise HTTPException(403, "Недостаточно прав для этого действия")

    def get_order(db, id_, user, lock=False):
        query = select(Order).where(Order.id == id_)
        if lock:
            query = query.with_for_update()
        order = db.scalar(query)
        if not order:
            raise HTTPException(404, "Наряд не найден")
        if user.role == "worker" and order.assignee_id != user.id:
            raise HTTPException(403, "Доступны только назначенные вам наряды")
        return order

    def resolve_assignee(db, assignee_id, brigade_id):
        if brigade_id:
            if not db.get(Brigade, brigade_id):
                raise HTTPException(422, "Бригада не найдена")
            workers = list(db.scalars(select(Employee).where(Employee.role == "worker", Employee.brigade_id == brigade_id, Employee.on_shift.is_(True)).order_by(Employee.id).with_for_update()))
            if not workers:
                raise HTTPException(422, "В бригаде нет работников на смене")
            counts = {w.id: db.scalar(select(func.count()).select_from(Order).where(Order.assignee_id == w.id, Order.status.notin_(TERMINAL))) for w in workers}
            return min(workers, key=lambda w: counts[w.id]).id
        person = db.get(Employee, assignee_id)
        if not person or person.role != "worker":
            raise HTTPException(422, "Исполнителем должен быть рабочий")
        if not person.on_shift:
            raise HTTPException(422, "Работник вне смены")
        return person.id

    async def changed(db, order):
        db.commit()
        await realtime.publish("orders.updated", order.id)
        await realtime.publish("notifications.updated")
        return order_dict(db, order, detail=True)

    def filtered(db, user, area_id=None, equipment_id=None, assignee_id=None, brigade_id=None, priority=None, status=None, search=None, from_date=None, to_date=None):
        query = select(Order)
        if user.role == "worker":
            query = query.where(Order.assignee_id == user.id)
        for key, value in [("area_id", area_id), ("equipment_id", equipment_id), ("assignee_id", assignee_id), ("brigade_id", brigade_id), ("priority", priority), ("status", status)]:
            if value is not None and value != "":
                if key == "status" and value not in STATUS:
                    raise HTTPException(422, "Неизвестный статус")
                if key == "priority" and value not in PRIORITY:
                    raise HTTPException(422, "Неизвестный приоритет")
                query = query.where(getattr(Order, key) == value)
        if search:
            safe_search = search.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
            needle = f"%{safe_search}%"
            query = query.where(Order.title.ilike(needle, escape="\\") | Order.number.ilike(needle, escape="\\") | Order.description.ilike(needle, escape="\\"))
        if from_date:
            query = query.where(Order.created_at >= parse_date(from_date))
        if to_date:
            query = query.where(Order.created_at <= parse_date(to_date, end=True))
        if from_date and to_date and parse_date(from_date) > parse_date(to_date, end=True):
            raise HTTPException(422, "Начало периода должно быть раньше окончания")
        return query

    @app.get("/api/health")
    @app.get("/health")
    def health(db: DB):
        db.execute(sql_text("SELECT 1"))
        return {"status": "ok", "database": "connected", "ai": "service" if getattr(ai_provider, "mode", None) == "service" else "openai" if ai_provider else "stub", "native": "stub"}

    @app.post("/api/auth/login")
    def login(payload: Login, request: Request, db: DB):
        key = (request.client.host if request.client else "unknown", payload.login)
        now = utcnow().timestamp()
        queue = attempts[key]
        while queue and queue[0] < now - 60:
            queue.popleft()
        if len(queue) >= 10:
            raise HTTPException(429, "Слишком много попыток. Повторите через минуту.", headers={"Retry-After": "60"})
        queue.append(now)
        person = db.scalar(select(Employee).where(Employee.login == payload.login))
        if not person or not check_pin(payload.pin, person.pin_hash):
            raise HTTPException(401, "Неверный логин или PIN")
        queue.clear()
        token = secrets.token_urlsafe(40)
        db.add(AuthSession(token_hash=token_hash(token), employee_id=person.id, expires_at=utcnow() + timedelta(hours=12)))
        db.execute(delete(AuthSession).where(AuthSession.expires_at < utcnow()))
        db.commit()
        return {"token": token, "user": employee_dict(person)}

    @app.get("/api/auth/me")
    def me(user: User):
        return employee_dict(user)

    @app.post("/api/auth/logout")
    def logout(request: Request, db: DB, user: User):
        db.execute(delete(AuthSession).where(AuthSession.token_hash == token_hash(request.headers["Authorization"][7:])))
        db.commit()
        return {"ok": True}

    @app.get("/api/reference")
    def reference(db: DB, user: User):
        return {"areas": rows(db, Area, ["id", "name"]), "equipment": rows(db, Equipment, ["id", "name", "inventory_number", "area_id", "type", "criticality"]), "employees": [employee_dict(p) for p in db.scalars(select(Employee).order_by(Employee.id))], "brigades": rows(db, Brigade, ["id", "name"]), "fault_codes": rows(db, FaultCode, ["id", "code", "name"]), "materials": rows(db, Material, ["id", "name", "unit"]), "time_norms": rows(db, TimeNorm, ["id", "name", "hours"])}

    def require_ai_service(request: Request):
        expected = os.getenv("AI_SERVICE_TOKEN", "")
        authorization = request.headers.get("Authorization", "")
        supplied = authorization[7:] if authorization.startswith("Bearer ") else ""
        if not expected:
            raise HTTPException(503, "Сервисный доступ ИИ не настроен")
        if not secrets.compare_digest(supplied, expected):
            raise HTTPException(401, "Неверный токен ИИ-сервиса")

    @app.get("/api/ai-service/snapshot")
    def ai_service_snapshot(request: Request, db: DB):
        require_ai_service(request)
        orders = list(db.scalars(select(Order).order_by(Order.id).limit(5000)))
        if len(orders) >= 5000:
            raise HTTPException(409, "История превышает 5000 нарядов; полный снимок недоступен")
        return {**reference(db=db, user=None), "orders": [order_dict(db, order, detail=True) for order in orders]}

    @app.get("/api/ai-service/photos/{id_}")
    def ai_service_photo(id_: int, request: Request, db: DB):
        require_ai_service(request)
        photo = db.get(Photo, id_)
        if not photo:
            raise HTTPException(404, "Фото не найдено")
        return Response(photo.data, media_type="image/jpeg", headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"})

    reference_models = {"areas": Area, "equipment": Equipment, "employees": Employee, "brigades": Brigade, "fault_codes": FaultCode, "materials": Material, "time_norms": TimeNorm}

    def save_reference(collection, payload, id_, db, user):
        require_role(user, "admin")
        model = reference_models.get(collection)
        if not model:
            raise HTTPException(404, "Справочник не найден")
        allowed = {c.name for c in model.__table__.columns} - {"id", "pin_hash"}
        if model is Employee:
            allowed.add("pin")
        if not payload or set(payload) - allowed:
            raise HTTPException(422, "Неизвестные поля или пустое изменение")
        existing = db.get(model, id_) if id_ else None
        if id_ and existing is None:
            raise HTTPException(404, "Запись не найдена")
        clean = {}
        for key, value in payload.items():
            if key in ["name", "login", "code", "unit", "type", "inventory_number", "specialty", "criticality"]:
                max_length = model.__table__.columns[key].type.length
                if not isinstance(value, str) or not value.strip() or len(value) > max_length:
                    raise HTTPException(422, f"Некорректное поле: {key}")
                clean[key] = value.strip()
            elif key in ["area_id", "brigade_id"]:
                if value is None and key == "brigade_id":
                    clean[key] = None
                elif isinstance(value, bool) or not isinstance(value, int) or not db.get(Area if key == "area_id" else Brigade, value):
                    raise HTTPException(422, f"Не найдена связь: {key}")
                else:
                    clean[key] = value
            elif key == "grade":
                if isinstance(value, bool) or not isinstance(value, int) or not 0 <= value <= 8:
                    raise HTTPException(422, "Разряд должен быть от 0 до 8")
                clean[key] = value
            elif key == "on_shift":
                if not isinstance(value, bool):
                    raise HTTPException(422, "on_shift должен быть boolean")
                clean[key] = value
            elif key == "role":
                if value not in ["worker", "master", "manager", "admin"]:
                    raise HTTPException(422, "Неизвестная роль")
                clean[key] = value
            elif key == "pin":
                if not isinstance(value, str) or not value.isdigit() or not 4 <= len(value) <= 12:
                    raise HTTPException(422, "PIN должен содержать от 4 до 12 цифр")
                clean["pin_hash"] = hash_pin(value)
            elif key == "hours":
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 < value <= 1000:
                    raise HTTPException(422, "Норматив должен быть больше 0 и не более 1000 часов")
                clean[key] = value
        required = {Area: {"name"}, Brigade: {"name"}, Equipment: {"name", "inventory_number", "area_id", "type"}, Employee: {"name", "login", "role", "pin_hash"}, FaultCode: {"code", "name"}, Material: {"name", "unit"}, TimeNorm: {"name", "hours"}}[model]
        if existing is None and not required.issubset(clean):
            raise HTTPException(422, "Заполните обязательные поля: " + ", ".join(sorted(required - set(clean))))
        if model is Employee and existing and existing.id == user.id and clean.get("role", existing.role) != "admin":
            raise HTTPException(422, "Нельзя изменить собственную роль администратора")
        obj = existing or model()
        for key, value in clean.items():
            setattr(obj, key, value)
        db.add(obj)
        try:
            db.commit()
        except IntegrityError:
            db.rollback()
            raise HTTPException(409, "Запись с таким значением уже существует")
        return employee_dict(obj) if model is Employee else {c.name: getattr(obj, c.name) for c in model.__table__.columns}

    @app.post("/api/reference/{collection}", status_code=201)
    def reference_create(collection: str, payload: dict, db: DB, user: User):
        return save_reference(collection, payload, None, db, user)

    @app.patch("/api/reference/{collection}/{id_}")
    def reference_update(collection: str, id_: int, payload: dict, db: DB, user: User):
        return save_reference(collection, payload, id_, db, user)

    @app.get("/api/employees")
    def employees(db: DB, user: User):
        all_orders = list(db.scalars(select(Order)))
        ratings = analytics(db, all_orders, utcnow() - timedelta(days=90), utcnow())["rankings"]
        rating_map = {r["id"]: r for r in ratings}
        result = []
        for person in db.scalars(select(Employee).where(Employee.role == "worker").order_by(Employee.id)):
            assigned = [o for o in all_orders if o.assignee_id == person.id and o.status not in EXECUTION_FINISHED and o.status != "rejected"]
            current = next((o for o in assigned if o.status in ["in_progress", "paused"]), None)
            rating = rating_map.get(person.id, {})
            result.append({**employee_dict(person), "status": "off_shift" if not person.on_shift else "busy" if current else "queued" if assigned else "free", "current_order": current.number if current else None, "queue_count": len([o for o in assigned if o is not current]), "rating": rating.get("score", 0), "completed_count": rating.get("closed_count", 0)})
        return result

    @app.get("/api/orders")
    def orders(db: DB, user: User, area_id: int | None = None, equipment_id: int | None = None, assignee_id: int | None = None, brigade_id: int | None = None, priority: str | None = None, status: str | None = None, search: str | None = Query(None, max_length=200), from_date: str | None = None, to_date: str | None = None, limit: int = Query(1000, ge=1, le=5000)):
        query = filtered(db, user, area_id, equipment_id, assignee_id, brigade_id, priority, status, search, from_date, to_date)
        refs = {"areas": {a.id: a for a in db.scalars(select(Area))}, "equipment": {e.id: e for e in db.scalars(select(Equipment))}, "employees": {p.id: p for p in db.scalars(select(Employee))}}
        return [order_dict(db, o, refs=refs) for o in db.scalars(query.order_by(Order.created_at.desc()).limit(limit))]

    @app.get("/api/orders/{id_}")
    def order_detail(id_: int, db: DB, user: User):
        return order_dict(db, get_order(db, id_, user), detail=True)

    @app.post("/api/orders", status_code=201)
    async def order_create(payload: OrderCreate, db: DB, user: User):
        require_role(user, "master", "admin")
        equipment = db.get(Equipment, payload.equipment_id)
        if not equipment or equipment.area_id != payload.area_id:
            raise HTTPException(422, "Оборудование не принадлежит выбранному участку")
        if payload.deadline <= utcnow():
            raise HTTPException(422, "Срок нового наряда должен быть в будущем")
        data = payload.model_dump()
        data["assignee_id"] = resolve_assignee(db, payload.assignee_id, payload.brigade_id)
        order = Order(**data, number=f"Н-{utcnow().year}-{secrets.token_hex(3).upper()}", status="issued", master_id=user.id)
        db.add(order)
        db.flush()
        audit(db, order, "issue", user.id, comment=payload.comment)
        notify(db, [order.assignee_id], "Вам назначен наряд", f"{order.number}: {order.title}", "assigned", order.id)
        return await changed(db, order)

    @app.patch("/api/orders/{id_}")
    async def order_update(id_: int, payload: OrderPatch, db: DB, user: User):
        require_role(user, "master", "admin")
        order = get_order(db, id_, user, lock=True)
        if order.status in TERMINAL:
            raise HTTPException(409, "Закрытый или отменённый наряд нельзя изменить")
        changes = payload.model_dump(exclude_unset=True)
        old_status = order.status
        audit_fields = []
        if "assignee_id" in changes or "brigade_id" in changes:
            if order.status in ["in_progress", "paused", "completed", "ai_review"]:
                raise HTTPException(409, "Переназначение доступно до начала работы или после возврата")
            order.assignee_id = resolve_assignee(db, payload.assignee_id, payload.brigade_id)
            order.brigade_id = payload.brigade_id
            order.status = "issued"
            notify(db, [order.assignee_id], "Наряд переназначен вам", order.number, "assigned", order.id)
        if payload.deadline is not None and payload.deadline <= utcnow():
            raise HTTPException(422, "Новый срок должен быть в будущем")
        for key, value in changes.items():
            if key not in ["assignee_id", "brigade_id"]:
                setattr(order, key, value)
            audit_fields.append(f"{key}={iso(value) if isinstance(value, datetime) else value}")
        audit(db, order, "edit", user.id, old_status, "; ".join(audit_fields))
        return await changed(db, order)

    @app.post("/api/orders/{id_}/transition")
    async def transition(id_: int, payload: Transition, db: DB, user: User):
        require_role(user, "worker", "master", "admin")
        order = get_order(db, id_, user, lock=True)
        action = payload.action
        if action in ["close", "rework", "cancel"]:
            require_role(user, "master", "admin")
        if action in ["reject", "pause", "rework", "cancel"] and not payload.reason:
            raise HTTPException(422, "Укажите причину действия")
        transitions = {"accept": ({"issued", "queued", "rework"}, "accepted"), "queue": ({"issued", "accepted", "rework"}, "queued"), "reject": ({"issued", "accepted", "queued"}, "rejected"), "start": ({"accepted", "queued", "rework"}, "in_progress"), "pause": ({"in_progress"}, "paused"), "resume": ({"paused"}, "in_progress"), "close": ({"ai_review"}, "closed"), "rework": ({"ai_review"}, "rework"), "cancel": (STATUS - TERMINAL, "cancelled")}
        allowed, target = transitions[action]
        if order.status not in allowed:
            raise HTTPException(409, f"Действие {action} недоступно для статуса {order.status}")
        if action in ["close", "rework"] and order.ai_review is None:
            raise HTTPException(409, "Дождитесь результата проверки ИИ или ручного fallback")
        if action in ["start", "resume"]:
            person = db.scalar(select(Employee).where(Employee.id == order.assignee_id).with_for_update())
            if not person.on_shift:
                raise HTTPException(409, "Исполнитель вне смены")
            busy = db.scalar(select(Order.id).where(Order.assignee_id == order.assignee_id, Order.id != order.id, Order.status == "in_progress"))
            if busy:
                raise HTTPException(409, "У исполнителя уже есть наряд в работе. Приостановите его или поставьте новый в очередь.")
            order.started_at = order.started_at or utcnow()
        old_status = order.status
        if action == "cancel" and order.work_type == "unplanned":
            order.downtime_minutes = downtime_minutes(order)
        order.status = target
        if action == "close":
            if payload.score is None:
                raise HTTPException(422, "Мастер должен поставить итоговую оценку от 1 до 5")
            suggested = (order.ai_review or {}).get("score")
            if (order.ai_review or {}).get("service_verdict") and (suggested is None or abs(payload.score - suggested) > 0.001) and not payload.reason.strip():
                raise HTTPException(422, "Укажите причину изменения оценки ИИ")
            order.score = payload.score
            order.closed_at = utcnow()
            order.ai_review = {**(order.ai_review or {}), "master_score": payload.score}
            assessment = db.scalar(select(AIAssessment).where(AIAssessment.order_id == order.id).order_by(AIAssessment.id.desc()).limit(1))
            if assessment:
                assessment.master_score = payload.score
        if action == "rework":
            order.completed_at = None
            order.score = None
        if action == "cancel":
            order.closed_at = utcnow()
        audit(db, order, action, user.id, old_status, payload.reason or payload.comment)
        notify(db, [order.assignee_id, order.master_id], "Статус наряда изменён", f"{order.number}: {old_status} → {target}", "status", order.id)
        return await changed(db, order)

    @app.post("/api/orders/{id_}/complete")
    async def complete(id_: int, payload: Completion, db: DB, user: User):
        require_role(user, "worker", "master", "admin")
        order = get_order(db, id_, user, lock=True)
        if order.status != "in_progress":
            raise HTTPException(409, "Завершить можно только наряд в работе")
        if not db.get(FaultCode, payload.fault_code_id):
            raise HTTPException(422, "Код неисправности не найден")
        if order.work_type == "unplanned" and not db.scalar(select(Photo.id).where(Photo.order_id == order.id, Photo.kind == "after")):
            raise HTTPException(422, "Для внеплановой работы добавьте фото после выполнения")
        materials = []
        for usage in payload.materials:
            material = db.get(Material, usage.material_id)
            if not material:
                raise HTTPException(422, f"Материал {usage.material_id} не найден")
            materials.append({"material_id": material.id, "name": material.name, "unit": material.unit, "quantity": usage.quantity})
            db.add(MaterialWriteoff(order_id=order.id, material_id=material.id, quantity=usage.quantity, author_id=user.id))
        previous_materials = (order.completion or {}).get("materials", [])
        accumulated = {m["material_id"]: dict(m) for m in previous_materials}
        for material in materials:
            if material["material_id"] in accumulated:
                accumulated[material["material_id"]]["quantity"] += material["quantity"]
            else:
                accumulated[material["material_id"]] = material
        order.completion = {**payload.model_dump(exclude={"materials"}), "materials": list(accumulated.values())}
        order.completed_at = utcnow()
        if order.work_type == "unplanned":
            order.downtime_minutes = round((order.completed_at - aware(order.created_at)).total_seconds() / 60, 1)
        order.status = "completed"
        completion_event = audit(db, order, "complete", user.id, "in_progress", payload.work_done)
        if ai_provider:
            db.flush()
            people = list(db.scalars(select(Employee)))
            fault = db.get(FaultCode, payload.fault_code_id)
            snapshot = {
                "order_number": order.number,
                "problem": redact_personnel(order.description or order.title, people),
                "work_done": redact_personnel(payload.work_done, people),
                "fault_code": fault.code,
                "fault_name": fault.name,
                "materials_this_attempt": materials,
                "normal_hours": order.normal_hours,
                "actual_hours": round((order.completed_at - aware(order.started_at)).total_seconds() / 3600, 2) if order.started_at else None,
                "deadline_met": aware(order.completed_at) <= aware(order.deadline),
                "submitted_by": user.id,
            }
            photo_ids = list(db.scalars(select(Photo.id).where(Photo.order_id == order.id).order_by(Photo.id)))
            db.add(AIReviewJob(order_id=order.id, completion_event_id=completion_event.id, snapshot=snapshot, photo_ids=photo_ids, status="pending", attempts=0, next_run_at=utcnow()))
            order.ai_review = None
            if getattr(ai_provider, "mode", None) == "service":
                order.status = "ai_review"
                audit(db, order, "ai_review_pending", user.id, "completed", "Проверка ИИ поставлена в очередь")
        else:
            order.ai_review = AIReviewStub.review(db, order)
            db.add(AIAssessment(order_id=order.id, **order.ai_review))
            order.status = "ai_review"
            audit(db, order, "ai_review", user.id, "completed", "Автоматическая проверка заглушкой ИИ. Ожидается решение мастера.")
            notify(db, [order.master_id], "Наряд ожидает приёмки", order.number, "review", order.id)
        return await changed(db, order)

    @app.post("/api/orders/{id_}/photos", status_code=201)
    async def upload_photo(id_: int, db: DB, user: User, file: UploadFile = File(...), kind: str = Form(...)):
        require_role(user, "worker", "master", "admin")
        order = get_order(db, id_, user, lock=True)
        if order.status in TERMINAL or order.status == "ai_review":
            raise HTTPException(409, "Фотографии нельзя менять после сдачи или закрытия наряда")
        if kind not in ["before", "after"]:
            raise HTTPException(422, "Тип фото: before или after")
        count = db.scalar(select(func.count()).select_from(Photo).where(Photo.order_id == id_, Photo.kind == kind))
        if count >= 5:
            raise HTTPException(422, "Можно загрузить не более 5 фотографий каждого типа")
        data = await file.read(10 * 1024 * 1024 + 1)
        await file.close()
        if len(data) > 10 * 1024 * 1024:
            raise HTTPException(413, "Размер фото не должен превышать 10 МБ")
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(data)) as original:
                    original.verify()
                with Image.open(io.BytesIO(data)) as original:
                    photo_image = ImageOps.exif_transpose(original).convert("RGB")
                    photo_image.thumbnail((1920, 1920))
                    output = io.BytesIO()
                    photo_image.save(output, format="JPEG", quality=82, optimize=True)
        except (UnidentifiedImageError, OSError, ValueError, SyntaxError, Image.DecompressionBombError, Image.DecompressionBombWarning):
            raise HTTPException(422, "Файл не является допустимым изображением")
        photo = Photo(order_id=id_, kind=kind, data=output.getvalue(), author_id=user.id)
        db.add(photo)
        audit(db, order, "photo", user.id, order.status, f"Добавлена фотография: {kind}")
        db.commit()
        await realtime.publish("orders.updated", id_)
        return photo_dict(db, photo)

    @app.get("/api/photos/{id_}")
    def get_photo(id_: int, db: DB, user: User):
        photo = db.get(Photo, id_)
        if not photo:
            raise HTTPException(404, "Фото не найдено")
        get_order(db, photo.order_id, user)
        return Response(photo.data, media_type="image/jpeg", headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"})

    @app.get("/api/dashboard")
    def dashboard(db: DB, user: User):
        orders = list(db.scalars(filtered(db, user)))
        start, label = shift_start()
        now = utcnow()
        active = [o for o in orders if o.status not in TERMINAL]
        completed = [o for o in orders if o.completed_at and aware(o.completed_at) >= start]
        scores = [o.score for o in orders if o.score is not None]
        return {"issued": sum(aware(o.created_at) >= start for o in orders), "completed": len(completed), "overdue": sum(aware(o.deadline) < now for o in active if o.status not in EXECUTION_FINISHED), "downtime_count": len({o.equipment_id for o in active if o.work_type == "unplanned" and o.status in ["in_progress", "paused", "rework"]}), "active": len(active), "total": len(orders), "avg_rating": round(sum(scores) / len(scores), 2) if scores else 0, "shift_label": label}

    @app.get("/api/notifications")
    def notifications(db: DB, user: User):
        return [{"id": n.id, "title": n.title, "message": n.message, "kind": n.kind, "order_id": n.order_id, "created_at": iso(n.created_at), "read": n.read} for n in db.scalars(select(Notification).where(Notification.employee_id == user.id).order_by(Notification.created_at.desc(), Notification.id.desc()).limit(200))]

    @app.post("/api/notifications/{id_}/read")
    async def mark_read(id_: int, db: DB, user: User):
        notification = db.get(Notification, id_)
        if not notification or notification.employee_id != user.id:
            raise HTTPException(404, "Уведомление не найдено")
        notification.read = True
        db.commit()
        await realtime.publish("notifications.updated")
        return {"ok": True}

    def analytics_data(db, user, days, from_date, to_date, area_id, equipment_id, assignee_id, brigade_id):
        start = parse_date(from_date) if from_date else utcnow() - timedelta(days=days)
        end = parse_date(to_date, end=True) if to_date else utcnow()
        if start > end:
            raise HTTPException(422, "Начало периода должно быть раньше окончания")
        if (end - start).days > 731:
            raise HTTPException(422, "Период не должен превышать 2 года")
        query = filtered(db, user, area_id, equipment_id, assignee_id, brigade_id).where(Order.created_at >= start, Order.created_at <= end)
        orders = list(db.scalars(query.order_by(Order.created_at.desc())))
        return orders, analytics(db, orders, start, end)

    @app.get("/api/analytics")
    def get_analytics(db: DB, user: User, days: int = Query(90, ge=1, le=731), from_date: str | None = None, to_date: str | None = None, area_id: int | None = None, equipment_id: int | None = None, assignee_id: int | None = None, brigade_id: int | None = None):
        return analytics_data(db, user, days, from_date, to_date, area_id, equipment_id, assignee_id, brigade_id)[1]

    @app.post("/api/ai/insights")
    async def ai_insights(db: DB, user: User, days: int = Query(90, ge=1, le=731), from_date: str | None = None, to_date: str | None = None, area_id: int | None = None, equipment_id: int | None = None, assignee_id: int | None = None, brigade_id: int | None = None):
        require_role(user, "master", "manager", "admin")
        if not ai_provider:
            raise HTTPException(503, "ИИ не настроен")
        if getattr(ai_provider, "mode", None) == "service":
            if any(value is not None for value in (equipment_id, assignee_id, brigade_id)):
                raise HTTPException(422, "ИИ-аналитика пока поддерживает фильтр только по участку")
            start = parse_date(from_date) if from_date else utcnow() - timedelta(days=days)
            end = parse_date(to_date, end=True) if to_date else utcnow()
            if end <= start or (end - start).days > 731:
                raise HTTPException(422, "Некорректный период аналитики")
            try:
                report = await ai_provider.get("/ai/analytics", {"start": iso(start), "end": iso(end), "area_id": area_id})
            except (httpx.HTTPError, ValueError):
                log.exception("AI service analytics failed")
                raise HTTPException(503, "ИИ-аналитика временно недоступна")
            findings = report.get("findings", [])
            if not isinstance(findings, list):
                raise HTTPException(502, "ИИ-сервис вернул некорректный отчёт")
            facts = [{"id": f"finding:{index}", "text": item["summary"]} for index, item in enumerate(findings)]
            finding_names = {"high_failure_equipment": "Частые поломки", "repeat_worker": "Повторный дефект",
                             "after_ppr": "Поломка после ППР", "shift_lateness": "Просрочки по сменам",
                             "material_overuse": "Аномальный расход", "unplanned_growth": "Рост внеплановых нарядов"}
            return {"summary": f"Проверено {report.get('orders_analyzed', 0)} нарядов; сигналов для проверки мастером: {len(findings)}.",
                    "insights": [{"title": finding_names.get(item["kind"], item["kind"]), "description": item["summary"],
                                  "recommendation": item["recommendation"], "fact_ids": [f"finding:{index}"],
                                  "order_ids": item.get("order_ids", [])} for index, item in enumerate(findings)],
                    "facts": facts, "model": "rules", "is_stub": False, "checked_without_llm": True}
        orders, _ = analytics_data(db, user, days, from_date, to_date, area_id, equipment_id, assignee_id, brigade_id)
        facts = evidence_facts(db, orders)
        db.rollback()
        try:
            result = await asyncio.to_thread(ai_provider.explain, facts, "Найди проверяемые закономерности и предложи действия мастеру", INSIGHTS_SCHEMA, "maintenance_insights")
            if not isinstance(result.get("summary"), str) or not isinstance(result.get("insights"), list):
                raise ValueError("Некорректный ответ модели")
            for insight in result["insights"]:
                supported_evidence(insight, facts)
                if not insight["fact_ids"]:
                    raise ValueError("Вывод без фактов")
            return {**result, "is_stub": False, "model": ai_provider.model, "facts": facts}
        except (httpx.HTTPError, ValueError, KeyError, TypeError):
            log.exception("AI insights failed")
            raise HTTPException(503, "ИИ-аналитика временно недоступна")

    @app.post("/api/ai/assistant")
    async def ai_assistant(payload: AssistantQuestion, db: DB, user: User):
        require_role(user, "master", "manager", "admin")
        if not ai_provider:
            raise HTTPException(503, "ИИ не настроен")
        if getattr(ai_provider, "mode", None) == "service":
            try:
                result = await ai_provider.post("/ai/assistant/ask", {"question": payload.question})
            except (httpx.HTTPError, ValueError):
                log.exception("AI service assistant failed")
                raise HTTPException(503, "ИИ-ассистент временно недоступен")
            facts = result.get("facts", {})
            return {"answer": result.get("answer", "Недостаточно данных"),
                    "facts": [{"id": key, "text": f"{key}: {value}"} for key, value in facts.items()],
                    "status": result.get("status"), "tool": result.get("tool"),
                    "is_stub": False, "checked_without_llm": True}
        orders = list(db.scalars(filtered(db, user)))
        people = list(db.scalars(select(Employee)))
        worker_names = {person.id: person.name for person in people if person.role == "worker"}
        facts = assistant_facts(db, orders)
        question = redact_personnel(payload.question, people)
        db.rollback()
        try:
            result = await asyncio.to_thread(ai_provider.explain, facts, question, ANSWER_SCHEMA, "master_answer")
            supported_evidence(result, facts)
            if not isinstance(result.get("answer"), str) or not result["answer"].strip():
                raise ValueError("Пустой ответ модели")
            result["answer"] = re.sub(r"сотрудник\s+#(\d+)", lambda match: worker_names.get(int(match.group(1)), match.group(0)), result["answer"], flags=re.IGNORECASE)
            return {**result, "is_stub": False, "model": ai_provider.model, "facts": [fact for fact in facts if fact["id"] in result["fact_ids"]]}
        except (httpx.HTTPError, ValueError, KeyError, TypeError):
            log.exception("AI assistant failed")
            raise HTTPException(503, "ИИ-ассистент временно недоступен")

    @app.post("/api/ai/order-hints")
    async def ai_order_hints(payload: OrderHintRequest, db: DB, user: User):
        require_role(user, "master", "admin")
        if not ai_provider:
            raise HTTPException(503, "ИИ не настроен")
        if getattr(ai_provider, "mode", None) == "service":
            try:
                result = await ai_provider.post("/ai/intake/text", {"phrase": payload.description})
            except (httpx.HTTPError, ValueError):
                log.exception("AI service order hints failed")
                raise HTTPException(503, "Подсказка ИИ временно недоступна")
            draft = result.get("draft") or {}
            return {"fault_code_id": draft.get("fault_code_id"), "time_norm_id": draft.get("time_norm_id"),
                    "explanation": "Подсказка по справочнику; проверьте шифр и норматив перед выдачей.",
                    "is_stub": True, "checked_without_llm": True}
        people = list(db.scalars(select(Employee)))
        fault_codes = list(db.scalars(select(FaultCode)))
        time_norms = list(db.scalars(select(TimeNorm)))
        fault_ids = {item.id for item in fault_codes}
        norm_ids = {item.id for item in time_norms}
        facts = [{"id": f"fault:{item.id}", "text": f"Шифр {item.code}: {item.name}"} for item in fault_codes]
        facts += [{"id": f"norm:{item.id}", "text": f"Норматив {item.name}: {item.hours} ч"} for item in time_norms]
        description = redact_personnel(payload.description, people)
        db.rollback()
        try:
            result = await asyncio.to_thread(ai_provider.explain, facts, description, HINT_SCHEMA, "order_hints")
            if result.get("fault_code_id") is not None and result["fault_code_id"] not in fault_ids:
                raise ValueError("Неизвестный шифр")
            if result.get("time_norm_id") is not None and result["time_norm_id"] not in norm_ids:
                raise ValueError("Неизвестный норматив")
            if not isinstance(result.get("explanation"), str):
                raise ValueError("Нет объяснения")
            return {**result, "is_stub": False, "model": ai_provider.model}
        except (httpx.HTTPError, ValueError, KeyError, TypeError):
            log.exception("AI order hints failed")
            raise HTTPException(503, "ИИ-подсказка временно недоступна")

    @app.get("/api/ai/reports/orders/{id_}")
    async def ai_order_report(id_: int, db: DB, user: User, format: str = Query("json", pattern="^(json|pdf|xlsx)$")):
        if getattr(ai_provider, "mode", None) != "service":
            raise HTTPException(503, "Отчёты отдельного ИИ-сервиса не подключены")
        get_order(db, id_, user)
        audience = "worker" if user.role == "worker" else "master"
        try:
            if format != "json":
                content = await ai_provider.download(f"/ai/reports/orders/{id_}", {"audience": audience, "format": format})
                media_type = "application/pdf" if format == "pdf" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                return Response(content, media_type=media_type,
                                headers={"Content-Disposition": f'attachment; filename="naryad-{id_}-ai-report.{format}"'})
            return await ai_provider.get(f"/ai/reports/orders/{id_}", {"audience": audience})
        except (httpx.HTTPError, ValueError):
            log.exception("AI service order report failed")
            raise HTTPException(503, "Отчёт ИИ временно недоступен")

    @app.get("/api/ai/reports/shift")
    async def ai_shift_report(start: datetime, end: datetime, user: User, format: str = Query("json", pattern="^(json|pdf|xlsx)$")):
        require_role(user, "master", "manager", "admin")
        if getattr(ai_provider, "mode", None) != "service":
            raise HTTPException(503, "Отчёты отдельного ИИ-сервиса не подключены")
        if end <= start:
            raise HTTPException(422, "Конец периода должен быть позже начала")
        try:
            params = {"start": start.isoformat(), "end": end.isoformat()}
            if format != "json":
                content = await ai_provider.download("/ai/reports/shift", {**params, "format": format})
                media_type = "application/pdf" if format == "pdf" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                return Response(content, media_type=media_type,
                                headers={"Content-Disposition": f'attachment; filename="naryad-shift-ai-report.{format}"'})
            return await ai_provider.get("/ai/reports/shift", params)
        except (httpx.HTTPError, ValueError):
            log.exception("AI service shift report failed")
            raise HTTPException(503, "Отчёт смены временно недоступен")

    @app.get("/api/ai/ratings")
    async def ai_ratings(start: datetime, end: datetime, user: User):
        if getattr(ai_provider, "mode", None) != "service":
            raise HTTPException(503, "Рейтинг отдельного ИИ-сервиса не подключён")
        if end <= start:
            raise HTTPException(422, "Конец периода должен быть позже начала")
        try:
            result = await ai_provider.get("/ai/ratings", {"start": start.isoformat(), "end": end.isoformat()})
        except (httpx.HTTPError, ValueError):
            log.exception("AI service ratings failed")
            raise HTTPException(503, "Рейтинг ИИ временно недоступен")
        if user.role == "worker":
            result["ratings"] = [item for item in result.get("ratings", []) if item.get("employee_id") == user.id]
        return result

    @app.get("/api/reports/export")
    def export_report(db: DB, user: User, days: int = Query(90, ge=1, le=731), from_date: str | None = None, to_date: str | None = None, area_id: int | None = None, equipment_id: int | None = None, assignee_id: int | None = None, brigade_id: int | None = None, format: str = Query("csv", pattern="^(csv|xlsx)$")):
        orders, report = analytics_data(db, user, days, from_date, to_date, area_id, equipment_id, assignee_id, brigade_id)
        headings = ["Номер", "Работа", "Участок", "Оборудование", "Исполнитель", "Тип", "Приоритет", "Статус", "Создан", "Срок", "Завершён", "Простой, мин", "Оценка"]
        output_rows = []
        for order in orders:
            o = order_dict(db, order)
            output_rows.append([safe_cell(o[k]) for k in ["number", "title", "area_name", "equipment_name", "assignee_name", "work_type", "priority", "status", "created_at", "deadline", "completed_at", "downtime_minutes", "score"]])
        if format == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Наряды"
            sheet.append(headings)
            for row in output_rows:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1D4B42")
            from openpyxl.utils import get_column_letter
            for column in range(1, len(headings) + 1):
                sheet.column_dimensions[get_column_letter(column)].width = 25 if column != 2 else 45
            summary_sheet = workbook.create_sheet("Сводка")
            for key, value in report["summary"].items():
                summary_sheet.append([key, value])
            summary_sheet.column_dimensions["A"].width = 30
            result = io.BytesIO()
            workbook.save(result)
            return Response(result.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="naryad-report.xlsx"'})
        result = io.StringIO()
        writer = csv.writer(result, delimiter=";")
        writer.writerow(headings)
        writer.writerows(output_rows)
        return Response(result.getvalue().encode("utf-8-sig"), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="naryad-report.csv"'})

    @app.get("/api/integrations")
    def integrations(user: User):
        service_mode = getattr(ai_provider, "mode", None) == "service"
        return {"ai": {"mode": "service" if service_mode else "openai" if ai_provider else "stub", "status": "active" if ai_provider else "demo", "description": "Проверка через отдельный ai_service; итоговое решение принимает мастер." if service_mode else "Проверка текста и фото через OpenAI; итоговое решение принимает мастер." if ai_provider else "Детерминированная заглушка. Реальные LLM и компьютерное зрение не подключены. Итоговое решение принимает мастер."}, "native": {"mode": "stub", "status": "demo", "description": "Контракт мобильного приложения и push-адаптер. События сохраняются в БД; отправки на устройства нет."}, "realtime": {"mode": "websocket", "status": "active", "description": "Авторизованный WebSocket и резервный опрос каждые 5 секунд."}}

    @app.websocket("/api/ws")
    @app.websocket("/ws")
    async def websocket(ws: WebSocket, token: str = ""):
        with sessions() as db:
            try:
                lookup_user(db, token)
            except HTTPException:
                await ws.close(code=1008)
                return
        await ws.accept()
        realtime.clients.add(ws)
        await ws.send_json({"type": "connected"})
        try:
            while True:
                # An idle connection must neither expire merely for being idle nor
                # retain access indefinitely after its authentication is revoked.
                try:
                    await asyncio.wait_for(ws.receive_text(), timeout=WS_AUTH_RECHECK_SECONDS)
                except asyncio.TimeoutError:
                    pass
                with sessions() as db:
                    lookup_user(db, token)
        except (WebSocketDisconnect, HTTPException):
            with suppress(Exception):
                await ws.close(code=1008)
        finally:
            realtime.clients.discard(ws)

    return app


def rows(db, model, fields):
    return [{key: getattr(obj, key) for key in fields} for obj in db.scalars(select(model).order_by(model.id))]


def parse_date(value, end=False):
    try:
        if len(value) == 10:
            day = date.fromisoformat(value)
            result = datetime.combine(day, datetime.min.time(), ZoneInfo("Asia/Almaty"))
            if end:
                result += timedelta(days=1, microseconds=-1)
            return result.astimezone(timezone.utc)
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return aware(parsed).astimezone(timezone.utc)
    except (ValueError, TypeError):
        raise HTTPException(422, "Дата должна быть в формате YYYY-MM-DD или ISO 8601")


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return "" if value is None else value


app = create_app()
